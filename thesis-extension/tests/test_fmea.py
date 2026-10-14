#!/usr/bin/env python3
"""Offline tests for FMEA coverage: the get_unmitigated_failure_modes tool, the FMEA
extractor's missing-value handling and provenance, and the link-proposal script.
No Neo4j is used.

Run from the repository root:
    python3 -m unittest discover -s thesis-extension/tests -v
"""

import contextlib
import csv
import importlib.util
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

TESTS_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS_DIR))
import test_inference_review as tir  # noqa: E402

sr8 = tir.sr8
EXT_DIR = TESTS_DIR.parent
EXTRACTION = EXT_DIR.parent / "thesis-extraction"
WORKBOOK = EXTRACTION / "Brake_System_Failure_Mode_Data.xlsx"


def _load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class FakeFmeaGraph:
    def __init__(self, rows, checked):
        self.rows, self.checked, self.queries = rows, checked, []

    def run(self, query, **params):
        self.queries.append((query, params))
        if "RETURN count(DISTINCT f) AS checked" in query:
            return tir.FakeResult([{"checked": self.checked}])
        return tir.FakeResult(self.rows)


class ToolTests(unittest.TestCase):
    def call(self, graph, keyword=""):
        saved = sr8._neo4j_driver
        sr8._neo4j_driver = lambda: tir.FakeDriver(graph)
        try:
            return json.loads(sr8.get_unmitigated_failure_modes(keyword))
        finally:
            sr8._neo4j_driver = saved

    ROW = {"requirement": "Brake wear adjustment", "component": "BrakeDiscs",
           "failure_mode_id": "FM-BrakeDiscs-WEAR_ADJUSTMENT_FAILURE", "failure_mode": "Wear adjustment failure",
           "effects": ["Brake clearance increases"], "causes": ["Adjuster mechanism inadequate"],
           "fm_file": "Brake_System_Failure_Mode_Data.xlsx", "fm_sheet": "Failure Mode Data", "fm_rows": [3],
           "link_file": "Brake_System_Failure_Mode_Data.xlsx", "link_sheet": "Requirement Mapping", "link_row": 2}

    def test_query_uses_the_extractors_names_and_the_no_mitigation_filter(self):
        graph = FakeFmeaGraph([], checked=20)
        self.call(graph, "  Wear ")
        query, params = graph.queries[0]
        self.assertIn("MATCH (r:Requirement)<-[s:satisfiesRequirement]-(c:Component)-[:hasFailureMode]->(f:FailureMode)", query)
        self.assertIn("AND NOT (f)-[:mitigatedBy]->(:ControlAction)", query)
        self.assertIn('($kw = "" OR toLower(r.hasName) CONTAINS toLower($kw))', query)
        for expr in ("r.hasName AS requirement", "c.hasID AS component", "f.hasName AS failure_mode",
                     "[(f)-[:causes]->(e:Effect) | e.hasName] AS effects"):
            self.assertIn(expr, query)
        self.assertEqual(params, {"kw": "Wear"})

    def test_result_has_provenance_pointer(self):
        out = self.call(FakeFmeaGraph([self.ROW], checked=20))
        self.assertEqual((out["checked_failure_modes"], out["unmitigated_count"]), (20, 1))
        prov = out["unmitigated_failure_modes"][0]["provenance"]
        self.assertEqual(prov["failure_mode"], {"workbook": "Brake_System_Failure_Mode_Data.xlsx",
                                                "sheet": "Failure Mode Data", "rows": [3]})
        self.assertEqual(prov["requirement_link"], {"workbook": "Brake_System_Failure_Mode_Data.xlsx",
                                                    "sheet": "Requirement Mapping", "row": 2})

    def test_graphs_without_stored_rows_still_point_to_the_sheet(self):
        row = dict(self.ROW, fm_file=None, fm_sheet=None, fm_rows=None, link_file=None, link_sheet=None, link_row=None)
        prov = self.call(FakeFmeaGraph([row], checked=1))["unmitigated_failure_modes"][0]["provenance"]
        self.assertEqual(prov["failure_mode"], {"workbook": sr8.FMEA_WORKBOOK, "sheet": "Failure Mode Data", "rows": None})
        self.assertEqual(prov["requirement_link"]["sheet"], "Requirement Mapping")

    def test_tool_never_links_fmea_and_adr_requirements(self):
        source = tir.SR8_PATH.read_text(encoding="utf-8")
        tool = source[source.index("def get_unmitigated_failure_modes"):source.index("# ── Phase 4")]
        self.assertNotIn("MERGE", tool)
        self.assertNotIn("CREATE", tool)
        self.assertNotIn("ExtractedEssentialRequirement)", tool)


class ExtractorTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.x = _load("extract_failure_modes_under_test", EXTRACTION / "Scripts" / "extract_failure_modes.py")

    def test_placeholders_are_missing(self):
        for value in ("", "  ", "None", "n/a", "NA", "-", "TBC", None):
            self.assertFalse(self.x.present(value), repr(value))
        self.assertTrue(self.x.present("provide_inspection_holes"))

    def test_dry_run_on_the_real_workbook(self):
        with contextlib.redirect_stdout(io.StringIO()):
            counts = self.x.extract_failure_modes(str(WORKBOOK), driver=None, dry_run=True)
        self.assertEqual((counts["components"], counts["failure_modes"], counts["distinct_failure_modes"],
                          counts["without_control_action"], counts["requirement_links"]), (15, 22, 20, 0, 23))

    def test_missing_control_action_creates_no_mitigation_and_rows_are_recorded(self):
        calls = []

        class Session:
            def run(self, query, **params):
                calls.append((query, params))

            def __enter__(self):
                return self

            def __exit__(self, *exc):
                return False

        class Driver:
            def session(self):
                return Session()

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "fmea.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Failure Mode Data"
            ws.append(["Component", "FailureMode", "Cause", "Effect", "ControlAction"])
            ws.append(["BrakeDiscs", "cracking", "thermal_load", "loss_of_braking", "inspect_discs"])
            ws.append(["BrakeDiscs", "warping", "overheating", "vibration", ""])
            ws.append(["BrakePads", "glazing", "n/a", "reduced_friction", "None"])
            req = wb.create_sheet("Requirement Mapping")
            req.append(["Component", "Requirement"])
            req.append(["BrakeDiscs", "Brake wear adjustment"])
            wb.save(path)
            with contextlib.redirect_stdout(io.StringIO()):
                counts = self.x.extract_failure_modes(str(path), driver=Driver(), dry_run=False)
        self.assertEqual(counts["without_control_action"], 2)
        mitigations = [p["fm_id"] for q, p in calls if "[:mitigatedBy]" in q]
        self.assertEqual(mitigations, ["FM-BrakeDiscs-CRACKING"])
        causes = [p["fm_id"] for q, p in calls if "[:causedBy]" in q]
        self.assertEqual(causes, ["FM-BrakeDiscs-CRACKING", "FM-BrakeDiscs-WARPING"])   # glazing cause is "n/a"
        self.assertFalse(any(p.get("node_id") in ("CA-", "CAUSE-", "EFFECT-") for _, p in calls))
        fm_rows = [(p["fm_id"], p["row"], p["sheet"]) for q, p in calls if "fm.sourceRows" in q]
        self.assertEqual(fm_rows, [("FM-BrakeDiscs-CRACKING", 2, "Failure Mode Data"),
                                   ("FM-BrakeDiscs-WARPING", 3, "Failure Mode Data"),
                                   ("FM-BrakePads-GLAZING", 4, "Failure Mode Data")])
        link = [p for q, p in calls if "rel.sourceRow = $row" in q]
        self.assertEqual([(p["requirement"], p["row"], p["sheet"]) for p in link],
                         [("Brake wear adjustment", 2, "Requirement Mapping")])


class ProposalScriptTests(unittest.TestCase):
    SCRIPT = EXT_DIR / "scripts" / "propose-fmea-links.py"

    def test_writes_review_csv_and_never_touches_the_graph(self):
        script = _load("propose_fmea_links_under_test", self.SCRIPT)
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "candidates.csv"
            with contextlib.redirect_stdout(io.StringIO()) as printed:
                script.main(["--output", str(out)])
            rows = list(csv.DictReader(open(out, newline="")))
        self.assertEqual(list(rows[0]), script.FIELDS)
        self.assertTrue(all(r["decision"] == r["reviewer"] == r["notes"] == "" for r in rows))
        per_req = {}
        for r in rows:
            per_req.setdefault(r["fmea_requirement"], []).append(float(r["name_overlap"]))
        self.assertTrue(all(len(v) <= 3 and v == sorted(v, reverse=True) and min(v) > 0 for v in per_req.values()))
        top = {r["fmea_requirement"]: r["candidate_req_id"] for r in rows if r["rank"] == "1"}
        self.assertEqual(top["Brake wear adjustment"], "REQ-31-038")   # "Brake Wear"
        self.assertIn("Nothing was written to the graph", printed.getvalue())
        source = self.SCRIPT.read_text(encoding="utf-8")
        for forbidden in ("import neo4j", "from neo4j", "GraphDatabase", "session.run", "MERGE", "CREATE"):
            self.assertNotIn(forbidden, source)

    def test_uses_the_shared_similarity_function(self):
        sim = _load("similarity_under_test", EXT_DIR / "similarity.py")
        self.assertEqual(sim.name_overlap("Brake wear adjustment", "Brake Wear"), 2 / 3)
        self.assertEqual(sr8._name_words("Linings of the Brake"), sim.name_words("Linings of the Brake"))
        self.assertIn('with_name("similarity.py")', tir.SR8_PATH.read_text(encoding="utf-8"))
        self.assertIn('"similarity.py"', self.SCRIPT.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
